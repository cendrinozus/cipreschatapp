import os, uuid
from pathlib import Path
from flask import current_app
from langchain_text_splitters import RecursiveCharacterTextSplitter
from app.services.embedder import add_chunks, delete_file_chunks

# ── Parsers ──────────────────────────────────────────────────────────────────

def _parse_pdf(path: str) -> str:
    from pypdf import PdfReader
    reader = PdfReader(path)
    return "\n".join(p.extract_text() or "" for p in reader.pages)

def _parse_docx(path: str) -> str:
    from docx import Document
    doc = Document(path)
    return "\n".join(p.text for p in doc.paragraphs)

def _parse_xlsx(path: str) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    lines = []
    for sheet in wb.worksheets:
        lines.append(f"[Feuille : {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            row_text = " | ".join(str(c) for c in row if c is not None)
            if row_text.strip():
                lines.append(row_text)
    return "\n".join(lines)

def _parse_txt(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()

PARSERS = {
    ".pdf":  _parse_pdf,
    ".docx": _parse_docx,
    ".xlsx": _parse_xlsx,
    ".txt":  _parse_txt,
    ".md":   _parse_txt,
}

# ── Ingestion principale ──────────────────────────────────────────────────────

def ingest_file(filepath: str, filename: str) -> dict:
    ext = Path(filename).suffix.lower()
    parser = PARSERS.get(ext)
    if not parser:
        raise ValueError(f"Format non supporté : {ext}")

    text = parser(filepath)
    if not text.strip():
        raise ValueError("Le document est vide ou illisible.")

    splitter = RecursiveCharacterTextSplitter(
        chunk_size    = current_app.config["MAX_CHUNK_SIZE"],
        chunk_overlap = current_app.config["CHUNK_OVERLAP"],
        separators    = ["\n\n", "\n", ".", " ", ""],
    )
    raw_chunks = splitter.split_text(text)

    chunks = []
    for i, chunk in enumerate(raw_chunks):
        chunks.append({
            "id":   f"{filename}__chunk_{i}_{uuid.uuid4().hex[:8]}",
            "text": chunk,
            "metadata": {
                "source":     filename,
                "chunk_index": i,
                "file_type":  ext.lstrip("."),
            },
        })

    add_chunks(chunks)
    return {"filename": filename, "chunks": len(chunks), "chars": len(text)}

def remove_file(filename: str):
    delete_file_chunks(filename)
    upload_folder = current_app.config["UPLOAD_FOLDER"]
    path = os.path.join(upload_folder, filename)
    if os.path.exists(path):
        os.remove(path)
